#!/usr/bin/env python3
"""
Builds the prospecting workbook for one HDB town (default Woodlands) from the
site's weekly HDB data:  python3 scripts/build_prospecting_sheet.py yishun out.xlsx

Tabs: How to use, Listings (blank, formula columns), Blocks, Block prices,
Street prices, Street list. The master copy lives in Serene's Google Drive
("North Prospecting – Master"); a monthly refresh pastes the data tabs over it.
Needs openpyxl. Not run by any workflow.
"""
import json, statistics as S, glob, os, re, sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import FormulaRule
from openpyxl.utils import get_column_letter as L
BASE=os.path.join(os.path.dirname(os.path.abspath(__file__)),'..')
R=os.path.join(BASE,'data')+'/'
TOWN=(sys.argv[1] if len(sys.argv)>1 else 'woodlands').lower()
OUT=sys.argv[2] if len(sys.argv)>2 else f'{TOWN.title()}_Prospecting_Data.xlsx'
m=json.load(open(R+'market-data.json')); last=m['hdb']['last_month']
def mb(ym,n):
    y,mm=int(ym[:4]),int(ym[5:7])-n
    while mm<=0: mm+=12;y-=1
    return f"{y:04d}-{mm:02d}"
c12=mb(last,12)
FT={"1":"1-room","2":"2-room","3":"3-room","4":"4-room","5":"5-room","E":"Executive","M":"Multi-gen"}
f=json.load(open(R+f'hdb/{TOWN}.json')); geo=json.load(open(R+f'geo/hdb/{TOWN}.json')); st=json.load(open(R+'geo/stations.json'))['stations']
pages={os.path.basename(p)[:-5] for p in glob.glob(os.path.join(BASE,'hdb',TOWN,'*.html'))}
slug=lambda s:re.sub(r'[^a-z0-9]+','-',s.lower()).strip('-')
title=lambda s:re.sub(r"(^|[\s/(])(\w)",lambda x:x.group(1)+x.group(2).upper(),s.lower())
med=lambda a:int(S.median(a)) if a else None
MON=["Jan","Feb","Mar","Apr","May","Jun","Jul","Aug","Sep","Oct","Nov","Dec"]
ym=lambda s:f"{MON[int(s[5:7])-1]} {s[:4]}"
year=int(last[:4])
blocks=[];prices=[];street_rows={}
for street,bl in f['streets'].items():
    for b,rs in bl.items():
        rs=sorted(rs,key=lambda r:r[0],reverse=True)
        r12=[r for r in rs if r[0]>c12]
        g=geo.get(f"{street}|{b}")
        types={}
        for r in rs: types.setdefault(r[1],[]).append(r)
        main=max(types,key=lambda k:len(types[k]))
        md=lambda ft: med([r[6] for r in r12 if r[1]==ft])
        p=f"blk-{slug(b)}-{slug(street)}"
        blocks.append([b,title(street),g[2] if g else "",len(rs),len(r12),FT[main],ym(rs[0][0]),rs[0][6],md("3"),md("4"),md("5"),md("E"),rs[0][5],99-(year-rs[0][5]),
            st[g[3][0][0]][0] if g else "", g[3][0][1] if g else None,
            f"https://sereneleeproperty.com/hdb/{TOWN}/{p}" if p in pages else "",
            f"https://sereneleeproperty.com/market-map?blk={b}&street={title(street).replace(' ','%20')}"])
        for ft,a in types.items():
            a12=[r for r in a if r[0]>c12]
            prices.append([f"{b}|{street}|{FT[ft]}".upper(),b,title(street),FT[ft],len(a),len(a12),med([r[6] for r in a]),med([r[6] for r in a12]),a[0][6],ym(a[0][0]),f"{a[0][2]:02d}-{a[0][2]+2:02d}"])
        for r in r12: street_rows.setdefault((street,FT[r[1]]),[]).append(r[6])
blocks.sort(key=lambda x:(-x[3],x[1],x[0]))
streets=sorted([[f"{s}|{t}".upper(),title(s),t,len(v),med(v)] for (s,t),v in street_rows.items()])
wb=Workbook()
H=Font(bold=True,color="FFFFFF"); HF=PatternFill("solid",fgColor="0B2A4A"); GOLD=PatternFill("solid",fgColor="F3E7BD")
def header(ws,cols,widths):
    ws.append(cols)
    for i,c in enumerate(cols,1):
        cell=ws.cell(1,i); cell.font=H; cell.fill=HF; cell.alignment=Alignment(wrap_text=True,vertical="center")
        ws.column_dimensions[L(i)].width=widths[i-1]
    ws.row_dimensions[1].height=34; ws.freeze_panes="A2"
# ---- How to use
hw=wb.active; hw.title="How to use"
lines=[(f"{TOWN.title()} prospecting sheet",True),("",False),
("1. Listings tab: add a PropertyGuru listing you want to follow. Fill the white columns: date listed, block, street, flat type, size, asking price, agent, agency, link.",False),
("   The grey columns fill themselves: asking $psf, what that flat type actually resold for in the same block (HDB records), the street median, how far the asking price is above or below, and days on the market.",False),
("   'Signal' flags listings above the block's resale median by more than 5% and listings older than 60 days: these owners are the most likely to switch agents or adjust price.",False),
("   Street names: pick from the dropdown (HDB spelling, e.g. 'Dr' for Drive). PropertyGuru's full spelling also works.",False),
("2. Blocks tab: all blocks with an HDB resale in the last 3 years, busiest first. Use it to choose blocks to farm (flyers, door-knocking). Links open your block page or the Market Map.",False),
("3. Block prices tab: the lookup table behind the Listings tab (every block and flat type). No need to edit.",False),
("",False),
(f"Data: HDB resale flat prices (data.gov.sg) to {ym(last)}, OneMap locations. Monthly refresh: open the new data file, and for each of Blocks, Block prices, Street prices and Street list, select all, copy, and paste over the same tab in the master (never delete those tabs, the formulas point to them).",False),
("Before contacting owners: respect exclusive agency agreements with other agents (CEA code of ethics) and check the DNC Registry for marketing messages.",False)]
for t,b in lines:
    hw.append([t]); 
    if b: hw.cell(hw.max_row,1).font=Font(bold=True,size=14,color="0B2A4A")
hw.column_dimensions["A"].width=150
# ---- Listings
ws=wb.create_sheet("Listings")
cols=["Date listed","Block","Street","Flat type","Size (sqft)","Asking price","Asking $psf","Block: last sale (same type)","Block: last sale month","Block median, same type (12 mths; 3 yrs if none)","Block resales, same type (3 yrs)","Street median, same type (12 mths)","Asking vs block median","Days listed","Signal","Agent","Agency","Listing link","Status","Next action","Notes","key","street key"]
header(ws,cols,[12,8,20,11,10,12,10,14,12,14,11,14,12,9,22,16,14,30,16,22,30,4,4])
NS='SUBSTITUTE(SUBSTITUTE(SUBSTITUTE(SUBSTITUTE(SUBSTITUTE(SUBSTITUTE(UPPER(TRIM(C{r})),"DRIVE","DR"),"AVENUE","AVE"),"STREET","ST"),"CRESCENT","CRES"),"ROAD","RD"),"CENTRAL","CTRL")'
from datetime import date
for r in range(2,202):
    ns=NS.format(r=r)
    row=[None]*len(cols)
    row[6]=f'=IF(AND(E{r}>0,F{r}>0),ROUND(F{r}/E{r},0),"")'
    row[21]=f'=IF(B{r}="","",UPPER(TRIM(B{r}))&"|"&{ns}&"|"&UPPER(D{r}))'
    row[22]=f'=IF(C{r}="","",{ns}&"|"&UPPER(D{r}))'
    row[7]=f"=IFERROR(VLOOKUP(V{r},'Block prices'!$A:$K,9,FALSE),\"\")"
    row[8]=f"=IFERROR(VLOOKUP(V{r},'Block prices'!$A:$K,10,FALSE),\"\")"
    row[9]=f"=IFERROR(IF(N(VLOOKUP(V{r},'Block prices'!$A:$K,8,FALSE))>0,VLOOKUP(V{r},'Block prices'!$A:$K,8,FALSE),VLOOKUP(V{r},'Block prices'!$A:$K,7,FALSE)),\"\")"
    row[10]=f"=IFERROR(VLOOKUP(V{r},'Block prices'!$A:$K,5,FALSE),\"\")"
    row[11]=f"=IFERROR(VLOOKUP(W{r},'Street prices'!$A:$E,5,FALSE),\"\")"
    row[12]=f'=IF(AND(F{r}>0,N(J{r})>0),F{r}/J{r}-1,"")'
    row[13]=f'=IF(A{r}="","",TODAY()-A{r})'
    row[14]=f'=IF(B{r}="","",TRIM(IF(AND(ISNUMBER(M{r}),M{r}>0.05),"Above block median ","")&IF(AND(ISNUMBER(N{r}),N{r}>60),"Stale (60+ days) ","")&IF(AND(ISNUMBER(M{r}),M{r}<=0.05,ISNUMBER(N{r}),N{r}<=60),"Fairly priced, fresh","")))'
    ws.append(row)
for r in range(2,202):
    ws[f"A{r}"].number_format="d mmm yyyy"
    for c in "FHJL": ws[f"{c}{r}"].number_format='"$"#,##0'
    ws[f"G{r}"].number_format='"$"#,##0'; ws[f"M{r}"].number_format='+0.0%;-0.0%;0.0%'
    for c in "GHIJKLMNO": ws[f"{c}{r}"].fill=PatternFill("solid",fgColor="F2F2F2")
ws.column_dimensions["V"].hidden=True; ws.column_dimensions["W"].hidden=True
dv=DataValidation(type="list",formula1='"2-room,3-room,4-room,5-room,Executive,Multi-gen"',allow_blank=True); ws.add_data_validation(dv); dv.add("D2:D201")
dv2=DataValidation(type="list",formula1='"New,Watching,Contacted agent (co-broke),Owner reached,Appointment set,Not interested,Sold"',allow_blank=True); ws.add_data_validation(dv2); dv2.add("S2:S201")
nst=len(sorted({s for s in [x[2] for x in prices]}))
dv3=DataValidation(type="list",formula1=f"='Street list'!$A$2:$A${nst+1}",allow_blank=True,showErrorMessage=False); ws.add_data_validation(dv3); dv3.add("C2:C201")
ws.conditional_formatting.add("O2:O201",FormulaRule(formula=['ISNUMBER(SEARCH("Above",O2))'],fill=PatternFill("solid",fgColor="F8D7D3")))
ws.conditional_formatting.add("N2:N201",FormulaRule(formula=['AND(ISNUMBER(N2),N2>60)'],fill=PatternFill("solid",fgColor="FCE8B2")))
ws.auto_filter.ref="A1:U201"
# ---- Blocks
bs=wb.create_sheet("Blocks")
bc=["Block","Street","Postal","Resales (3 yrs)","Resales (12 mths)","Most traded type","Last resale","Last price","3-room median (12 mths)","4-room median (12 mths)","5-room median (12 mths)","Executive median (12 mths)","Lease from","Lease left (yrs)","Nearest MRT","MRT distance (m)","Block page","Market Map","Farming status","Notes"]
header(bs,bc,[8,20,9,10,10,12,11,11,12,12,12,12,9,9,20,10,16,14,16,30])
for x in blocks:
    row=x[:16]+[None,None,None,None]; bs.append(row); r=bs.max_row
    if x[16]: bs.cell(r,17).value=f'=HYPERLINK("{x[16]}","Block page")'
    bs.cell(r,18).value=f'=HYPERLINK("{x[17]}","Check block")'
    for c in (8,9,10,11,12): bs.cell(r,c).number_format='"$"#,##0'
dv4=DataValidation(type="list",formula1='"Target,Flyered,Door-knocked,Follow up,Skip"',allow_blank=True); bs.add_data_validation(dv4); dv4.add(f"S2:S{bs.max_row}")
bs.auto_filter.ref=f"A1:T{bs.max_row}"
# ---- Block prices
ps=wb.create_sheet("Block prices")
header(ps,["Key","Block","Street","Flat type","Resales (3 yrs)","Resales (12 mths)","Median (3 yrs)","Median (12 mths)","Last price","Last sale","Last floors"],[30,8,20,11,10,10,12,12,12,11,10])
for p in sorted(prices): ps.append(p)
for r in range(2,ps.max_row+1):
    for c in (7,8,9): ps.cell(r,c).number_format='"$"#,##0'
ss=wb.create_sheet("Street prices")
header(ss,["Key","Street","Flat type","Resales (12 mths)","Median (12 mths)"],[30,22,11,10,12])
for s in streets: ss.append(s)
sl=wb.create_sheet("Street list"); sl.append([f"HDB street names ({TOWN.title()})"])
for s in sorted({x[2] for x in prices}): sl.append([s])
wb.move_sheet("Listings",offset=-1)
wb.active=0
out=OUT; wb.save(out)
print(len(blocks),len(prices),len(streets),os.path.getsize(out))
